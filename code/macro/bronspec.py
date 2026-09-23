"""
Specificatie van ALLE inputreeksen van de macro-analyse.

Voor elke reeks leggen we vast:
  - waar we ze OVERNEMEN (werkboek van Willem of RSVZ-werkboek: blad, rij, kolommen),
  - waar ze OORSPRONKELIJK vandaan komt (organisatie, bronbestand, blad, rij),
  - label, eenheid, jaren en eventuele opmerkingen.

Deze ene lijst wordt door drie scripts gebruikt:
  01_extract_raw.py     leest de waarden uit het werkboek -> data/raw/*.csv
                        en schrijft data/raw/codebook.csv
  02_verify_sources.py  controleert of die waarden gelijk zijn aan het
                        oorspronkelijke bronbestand
  (en het codeboek komt mee in de gegenereerde Excel)

Wie een reeks wil traceren, vindt hier dus in één regel de volledige keten:
bron -> werkboek -> CSV.
"""

from dataclasses import dataclass, field
from typing import Optional

from openpyxl.utils import get_column_letter


def kolommen(eerste_kolom: int, eerste_jaar: int, laatste_jaar: int) -> str:
    """Geef een leesbaar kolombereik, bv. 'B:V', voor een reeks jaren.

    eerste_kolom is een kolomnummer (1 = A, 2 = B, ...).
    """
    laatste_kolom = eerste_kolom + (laatste_jaar - eerste_jaar)
    return f"{get_column_letter(eerste_kolom)}:{get_column_letter(laatste_kolom)}"


@dataclass
class Reeks:
    csv: str                 # doelbestand in data/raw (zonder .csv)
    variabele: str           # kolomnaam in dat CSV-bestand
    label: str               # omschrijving in mensentaal
    eenheid: str
    jaren: tuple             # (eerste jaar, laatste jaar)

    # --- waar we de waarden overnemen ---
    werkboek: str            # 'willem' of 'rsvz'
    blad: str
    rij: int
    kol_eerste_jaar: int     # kolomnummer van het eerste jaar in dat blad

    # --- oorspronkelijke bron (voor traceerbaarheid en verificatie) ---
    bron_organisatie: str
    bron_bestand: str        # pad t.o.v. data/sources, of omschrijving als er geen bestand is
    bron_locatie: str        # blad/rij/kolommen in het bronbestand, in woorden
    geraadpleegd: str = ""
    opmerking: str = ""

    # Machineleesbare verwijzing voor 02_verify_sources.py.
    # Vorm: (soort, parameters). None = geen bronbestand beschikbaar.
    verificatie: Optional[tuple] = field(default=None)

    @property
    def werkboek_locatie(self) -> str:
        bestand = {
            "willem": "werkboeken/macro_analyses_willem_2026-09-16.xlsx",
            "rsvz": "rsvz/rsvz_samenstelling_zelfstandigen_2000-2024.xlsx",
            "fodsz_z": "fodsz/fodsz_zelfstandigen_2000-2025.xlsm (RECHTSTREEKS UIT BRON, niet in werkboek)",
            "fodsz_w": "fodsz/fodsz_werknemers_2000-2025.xlsm (RECHTSTREEKS UIT BRON, niet in werkboek)",
        }[self.werkboek]
        k = kolommen(self.kol_eerste_jaar, *self.jaren)
        return f"{bestand} :: '{self.blad}'!rij {self.rij}, kolommen {k}"


# ---------------------------------------------------------------------------
# Vaste stukjes tekst
# ---------------------------------------------------------------------------
FODSZ_Z = "fodsz/fodsz_zelfstandigen_2000-2025.xlsm"
FODSZ_W = "fodsz/fodsz_werknemers_2000-2025.xlsm"
FODSZ_ORG = "FOD Sociale Zekerheid, ontvangsten en uitgaven globaal beheer (persoonlijke communicatie, 18 april 2025)"
FODSZ_OPM = ("Bedragen in duizend euro, lopende rekeningen. 2000-2007 inclusief geneeskundige verzorging "
             "(zie fodsz_opmerkingen_bij_tabellen.docx). 2024-2025 zijn MONITOR_03.2025-ramingen, geen realisaties.")
KSZ_ORG = ("KSZ, Datawarehouse Arbeidsmarkt en Sociale Bescherming, nomenclatuur socio-economische positie, "
           "webtoepassing 'globale cijfers' (https://sami.ksz-bcss.fgov.be/samigc/homePage.xhtml)")
KSZ_BESTAND = "GEEN BRONBESTAND: manueel overgenomen uit de webtoepassing, 4e kwartaal van elk jaar"
KSZ_OPM = ("Er is geen export van de webtoepassing bewaard. Aanbevolen: de reeks opnieuw exporteren en "
           "als bestand in data/sources/ksz/ opslaan.")
NBB_BESTAND = "nbb/nbb_sociale_premies_2003-2023.xlsx"
NBB_ORG = "Nationale Bank van België, NBB.Stat, Ontvangen belastingen en werkelijke sociale premies, Totale overheid"
RSVZ_BESTAND = "rsvz/rsvz_samenstelling_zelfstandigen_2000-2024.xlsx"
RSVZ_ORG = "RSVZ statistiekendatabase (https://websta.rsvz-inasti.fgov.be/nl)"
STATBEL_BESTAND = "statbel/statbel_cpi_historiek_1920-2025.xls"
STATBEL_ORG = "Statbel, Consumptieprijsindex vanaf 1920 en gezondheidsindex vanaf 1994"


REEKSEN = [
    # =======================================================================
    # Prijsindexen
    # =======================================================================
    Reeks("prijsindex", "cpi_2013", "Consumptieprijsindex, jaargemiddelde", "index (2013 = 100)", (2003, 2024),
          "willem", "Julie BS - Cijfers", 2, 2,
          STATBEL_ORG, STATBEL_BESTAND, "blad 'general index', rij 'Jaar/Année' per jaar, kolom basis 2013",
          geraadpleegd="bestand 'conv-2025'",
          opmerking="Afgerond op 2 decimalen in het werkboek (Statbel publiceert meer decimalen).",
          verificatie=("statbel", "general index")),
    Reeks("prijsindex", "gezondheidsindex_2013", "Gezondheidsindex, jaargemiddelde", "index (2013 = 100)", (2003, 2024),
          "willem", "Julie BS - gezondheidsindex", 2, 2,
          STATBEL_ORG, STATBEL_BESTAND, "blad 'health index', rij 'Année/Jaar' per jaar, kolom basis 2013",
          opmerking="Enkel gebruikt voor de robuustheidscheck in voetnoot 12.",
          verificatie=("statbel", "health index")),

    # =======================================================================
    # KSZ-populaties (bouwstenen van Z1, Z2, W1, W2)
    # =======================================================================
    Reeks("ksz_populatie", "n11", "n11 Werkend in loondienst", "personen (4e kwartaal)", (2003, 2023),
          "willem", "Oplijsting (sub)categorieën KSZ", 2, 3, KSZ_ORG, KSZ_BESTAND, "positie n11",
          geraadpleegd="22 december 2025 (Figuur 1) / 21 januari 2026 (voetnoot 14)", opmerking=KSZ_OPM),
    Reeks("ksz_populatie", "n12", "n12 Werkend als zelfstandige", "personen (4e kwartaal)", (2003, 2023),
          "willem", "Oplijsting (sub)categorieën KSZ", 3, 3, KSZ_ORG, KSZ_BESTAND, "positie n12",
          opmerking=KSZ_OPM),
    Reeks("ksz_populatie", "n13", "n13 Werkend als helper", "personen (4e kwartaal)", (2003, 2023),
          "willem", "Oplijsting (sub)categorieën KSZ", 4, 3, KSZ_ORG, KSZ_BESTAND, "positie n13",
          opmerking=KSZ_OPM),
    Reeks("ksz_populatie", "n141", "n141 Loondienst en zelfstandige/helper, voornaamste job in loondienst",
          "personen (4e kwartaal)", (2003, 2023),
          "willem", "Oplijsting (sub)categorieën KSZ", 5, 3, KSZ_ORG, KSZ_BESTAND, "positie n141",
          opmerking=KSZ_OPM),
    Reeks("ksz_populatie", "n142", "n142 Loondienst en zelfstandige/helper, voornaamste job als zelfstandige",
          "personen (4e kwartaal)", (2003, 2023),
          "willem", "Oplijsting (sub)categorieën KSZ", 6, 3, KSZ_ORG, KSZ_BESTAND, "positie n142",
          opmerking=KSZ_OPM),
    Reeks("ksz_populatie", "n143", "n143 Loondienst en zelfstandige/helper, voornaamste job als helper",
          "personen (4e kwartaal)", (2003, 2023),
          "willem", "Oplijsting (sub)categorieën KSZ", 7, 3, KSZ_ORG, KSZ_BESTAND, "positie n143",
          opmerking=KSZ_OPM),

    # =======================================================================
    # FOD SZ, zelfstandigen: bedragen (beroepssolidariteit)
    # =======================================================================
    Reeks("fodsz_bedragen", "z_bijdragen", "Zelfstandigen: bijdragen (code 101)", "duizend euro, nominaal", (2003, 2025),
          "willem", "Julie BS - Cijfers", 8, 2, FODSZ_ORG, FODSZ_Z,
          "blad 'Zelfstandigen - Indépendants', rij 23 (code 101), kolommen D:AC = 2000-2025", opmerking=FODSZ_OPM,
          verificatie=("fodsz_z", 23)),
    Reeks("fodsz_bedragen", "z_gewone_bijdragen", "Zelfstandigen: gewone bijdragen (code 101.1)", "duizend euro, nominaal",
          (2003, 2025), "willem", "Julie BS - Cijfers", 63, 2, FODSZ_ORG, FODSZ_Z,
          "blad 'Zelfstandigen - Indépendants', rij 24 (code 101.1)", opmerking=FODSZ_OPM,
          verificatie=("fodsz_z", 24)),
    Reeks("fodsz_bedragen", "z_vennootschapsbijdragen", "Zelfstandigen: vennootschapsbijdragen (code 101.5)",
          "duizend euro, nominaal", (2003, 2025), "willem", "Julie BS - Cijfers", 79, 2, FODSZ_ORG, FODSZ_Z,
          "blad 'Zelfstandigen - Indépendants', rij 25 (code 101.5)", opmerking=FODSZ_OPM,
          verificatie=("fodsz_z", 25)),
    Reeks("fodsz_bedragen", "z_uitgaven", "Zelfstandigen: totaal lopende uitgaven (code 110.1)", "duizend euro, nominaal",
          (2003, 2025), "willem", "Julie BS - Cijfers", 9, 2, FODSZ_ORG, FODSZ_Z,
          "blad 'Zelfstandigen - Indépendants', rij 84 (code 110.1)", opmerking=FODSZ_OPM,
          verificatie=("fodsz_z", 84)),

    # =======================================================================
    # FOD SZ, werknemers: bedragen (beroepssolidariteit)
    # =======================================================================
    Reeks("fodsz_bedragen", "w_bijdragen", "Werknemers: bijdragen (code 101)", "duizend euro, nominaal", (2003, 2025),
          "willem", "Julie BS - Cijfers", 35, 2, FODSZ_ORG, FODSZ_W,
          "blad 'Werknemers - Salariés', rij 17 (code 101), kolommen D:AC = 2000-2025",
          opmerking=FODSZ_OPM + " Bevat werkgevers- en werknemersbijdragen samen.",
          verificatie=("fodsz_w", 17)),
    Reeks("fodsz_bedragen", "w_uitgaven", "Werknemers: totaal lopende uitgaven (code 110.1)", "duizend euro, nominaal",
          (2003, 2025), "willem", "Julie BS - Cijfers", 36, 2, FODSZ_ORG, FODSZ_W,
          "blad 'Werknemers - Salariés', rij 90 (code 110.1)", opmerking=FODSZ_OPM,
          verificatie=("fodsz_w", 90)),

    # =======================================================================
    # FOD SZ: aandelen in de totale lopende ontvangsten (nationale solidariteit)
    # In het werkboek staan enkel de aandelen (hard gecodeerd), niet de bedragen.
    # =======================================================================
    Reeks("fodsz_aandelen", "z_aandeel_bijdragen", "Zelfstandigen: aandeel bijdragen in totale lopende ontvangsten",
          "fractie (0-1)", (2000, 2024), "willem", "Julie NS - cijfers", 3, 2, FODSZ_ORG, FODSZ_Z,
          "blad 'Zelfstandigen - Indépendants', rij 23 / rij 57 (in het bronbestand ook uitgerekend in rij 58)",
          opmerking=FODSZ_OPM, verificatie=("fodsz_z_aandeel", 23, 57)),
    Reeks("fodsz_aandelen", "z_aandeel_toelagen", "Zelfstandigen: aandeel toelagen van publieke overheden (code 102)",
          "fractie (0-1)", (2000, 2024), "willem", "Julie NS - cijfers", 4, 2, FODSZ_ORG, FODSZ_Z,
          "rij 30 / rij 57 (bronbestand rij 59)", opmerking=FODSZ_OPM, verificatie=("fodsz_z_aandeel", 30, 57)),
    Reeks("fodsz_aandelen", "z_aandeel_alternatieve", "Zelfstandigen: aandeel alternatieve financiering (code 103)",
          "fractie (0-1)", (2000, 2024), "willem", "Julie NS - cijfers", 5, 2, FODSZ_ORG, FODSZ_Z,
          "rij 35 / rij 57 (bronbestand rij 60)", opmerking=FODSZ_OPM, verificatie=("fodsz_z_aandeel", 35, 57)),
    Reeks("fodsz_aandelen", "z_aandeel_overige", "Zelfstandigen: aandeel overige financiering (codes 104-107)",
          "fractie (0-1)", (2000, 2024), "willem", "Julie NS - cijfers", 6, 2, FODSZ_ORG, FODSZ_Z,
          "rijen 51-54 / rij 57 (bronbestand rij 64)", opmerking=FODSZ_OPM,
          verificatie=("fodsz_z_aandeel", (51, 52, 53, 54), 57)),
    Reeks("fodsz_aandelen", "w_aandeel_bijdragen", "Werknemers: aandeel bijdragen in totale lopende ontvangsten",
          "fractie (0-1)", (2000, 2024), "willem", "Julie NS - cijfers", 17, 2, FODSZ_ORG, FODSZ_W,
          "blad 'Werknemers - Salariés', rij 17 / rij 58 (bronbestand rij 60)", opmerking=FODSZ_OPM,
          verificatie=("fodsz_w_aandeel", 17, 58)),
    Reeks("fodsz_aandelen", "w_aandeel_toelagen", "Werknemers: aandeel toelagen van publieke overheden (code 102)",
          "fractie (0-1)", (2000, 2024), "willem", "Julie NS - cijfers", 18, 2, FODSZ_ORG, FODSZ_W,
          "rij 23 / rij 58 (bronbestand rij 61)", opmerking=FODSZ_OPM, verificatie=("fodsz_w_aandeel", 23, 58)),
    Reeks("fodsz_aandelen", "w_aandeel_alternatieve", "Werknemers: aandeel alternatieve financiering (code 103)",
          "fractie (0-1)", (2000, 2024), "willem", "Julie NS - cijfers", 19, 2, FODSZ_ORG, FODSZ_W,
          "rij 31 / rij 58 (bronbestand rij 62)", opmerking=FODSZ_OPM, verificatie=("fodsz_w_aandeel", 31, 58)),
    Reeks("fodsz_aandelen", "w_aandeel_overige", "Werknemers: aandeel overige financiering (codes 104-107)",
          "fractie (0-1)", (2000, 2024), "willem", "Julie NS - cijfers", 20, 2, FODSZ_ORG, FODSZ_W,
          "rijen 46, 53, 54, 55 / rij 58 (bronbestand rij 66)", opmerking=FODSZ_OPM,
          verificatie=("fodsz_w_aandeel", (46, 53, 54, 55), 58)),

    # =======================================================================
    # NBB: opsplitsing werknemers- en werkgeversbijdragen (Bijlage 5)
    # =======================================================================
    Reeks("nbb_bijdragen", "werknemersbijdragen_d613", "Werkelijke sociale premies t.l.v. de werknemers (D.613)",
          "duizend euro, nominaal", (2003, 2023), "willem", "Julie BS - Cijfers", 92, 2, NBB_ORG, NBB_BESTAND,
          "blad 'Table', rij 47 (duizend euro), kolommen D:X = 2003-2023",
          geraadpleegd="NBB.Stat laatst bijgewerkt 5 november 2025",
          opmerking=("Nationale rekeningen, totale overheid: breder aggregaat dan het FOD SZ-globaal beheer "
                     "(o.a. ook statutaire ambtenaren). Zie leesnota G3."),
          verificatie=("nbb", 47)),
    Reeks("nbb_bijdragen", "werkgeversbijdragen_d611", "Werkelijke sociale premies t.l.v. de werkgevers (D.611)",
          "duizend euro, nominaal", (2003, 2023), "willem", "Julie BS - Cijfers", 93, 2, NBB_ORG, NBB_BESTAND,
          "blad 'Table', rij 38 (duizend euro)", geraadpleegd="NBB.Stat laatst bijgewerkt 5 november 2025",
          opmerking="Zie opmerking bij D.613.", verificatie=("nbb", 38)),

    # =======================================================================
    # RSVZ: samenstelling zelfstandigenpopulatie (Figuur 2, Figuur 5)
    # =======================================================================
    Reeks("rsvz_samenstelling", "vennootschappen", "Aantal actieve vennootschappen", "aantal", (2000, 2024),
          "rsvz", "Cijfers", 3, 2, RSVZ_ORG, "Geen export bewaard; waarden in het RSVZ-werkboek (blad 'Cijfers')",
          "ondernemingsvorm: vennootschappen", geraadpleegd="24 oktober 2025",
          opmerking="Ook in het werkboek van Willem ('Julie BS - Cijfers' rij 78, 2003-2023); daar gecontroleerd.",
          verificatie=("willem_rij", "Julie BS - Cijfers", 78)),
    Reeks("rsvz_samenstelling", "zelfstandigen", "Aantal zelfstandigen (hoedanigheid)", "aantal", (2000, 2024),
          "rsvz", "Cijfers", 8, 2, RSVZ_ORG, "Geen export bewaard; waarden in het RSVZ-werkboek", "hoedanigheid",
          geraadpleegd="24 oktober 2025"),
    Reeks("rsvz_samenstelling", "helpers", "Aantal helpers (hoedanigheid)", "aantal", (2000, 2024),
          "rsvz", "Cijfers", 9, 2, RSVZ_ORG, "Geen export bewaard; waarden in het RSVZ-werkboek", "hoedanigheid",
          geraadpleegd="24 oktober 2025",
          opmerking="Breuk 2002-2003: verplichtstelling ministatuut meewerkende echtgenoten (leesnota G5)."),
    Reeks("rsvz_samenstelling", "hoofdberoep", "Aard van de bezigheid: hoofdberoep", "aantal", (2000, 2024),
          "rsvz", "Cijfers", 17, 2, RSVZ_ORG, "Geen export bewaard; waarden in het RSVZ-werkboek", "aard bezigheid",
          geraadpleegd="24 oktober 2025"),
    Reeks("rsvz_samenstelling", "bijberoep", "Aard van de bezigheid: bijberoep", "aantal", (2000, 2024),
          "rsvz", "Cijfers", 18, 2, RSVZ_ORG, "Geen export bewaard; waarden in het RSVZ-werkboek", "aard bezigheid",
          geraadpleegd="24 oktober 2025"),
    Reeks("rsvz_samenstelling", "actief_na_pensioen", "Aard van de bezigheid: actief na pensioen", "aantal",
          (2000, 2024), "rsvz", "Cijfers", 19, 2, RSVZ_ORG,
          "Geen export bewaard; waarden in het RSVZ-werkboek", "aard bezigheid", geraadpleegd="24 oktober 2025"),
] + [
    Reeks("rsvz_samenstelling", var, f"Bedrijfstak: {lab}", "aantal", (2000, 2024), "rsvz", "Cijfers", rij, 2,
          RSVZ_ORG, "Geen export bewaard; waarden in het RSVZ-werkboek", "bedrijfstak (beroepscodes, zie voetnoot 2)",
          geraadpleegd="24 oktober 2025")
    for var, lab, rij in [
        ("landbouw", "landbouw", 29), ("visserij", "visserij", 30), ("nijverheid", "nijverheid", 31),
        ("handel", "handel", 32), ("vrije_beroepen", "vrije beroepen", 33), ("diensten", "diensten", 34),
        ("diversen", "diversen", 35),
    ]
]

# ===========================================================================
# Aanvullende FOD SZ-reeksen voor de interpretatie in §6.1.
# Deze staan NIET in het werkboek van Willem en worden rechtstreeks uit het
# originele FOD SZ-bestand gelezen (kolom D = 2000). Een '-' in de bron = 0.
# ===========================================================================
_BLAD_Z, _BLAD_W = "Zelfstandigen - Indépendants", "Werknemers - Salariés"
REEKSEN += [
    Reeks("fodsz_componenten", var, lab, "duizend euro, nominaal", (2000, 2025), wb, blad, rij, 4,
          FODSZ_ORG, bron, f"blad '{blad}', rij {rij} ({code})", opmerking=FODSZ_OPM + opm)
    for var, lab, wb, blad, rij, code, bron, opm in [
        ("z_pensioenen", "Zelfstandigen: RSVZ-pensioenen", "fodsz_z", _BLAD_Z, 71, "111 PE", FODSZ_Z, ""),
        ("z_gezinsbijslag", "Zelfstandigen: RSVZ-gezinsbijslag", "fodsz_z", _BLAD_Z, 72, "111 KB", FODSZ_Z,
         " Verdwijnt vanaf 2015 (6de staatshervorming)."),
        ("z_naar_riziv", "Zelfstandigen: overdrachten naar RIZIV-geneeskundige verzorging", "fodsz_z", _BLAD_Z, 79,
         "114.2", FODSZ_Z, " Tot 2007 zit geneeskundige verzorging in de prestaties zelf."),
        ("z_toelagen_bedrag", "Zelfstandigen: toelagen van publieke overheden", "fodsz_z", _BLAD_Z, 30, "102",
         FODSZ_Z, ""),
        ("z_alternatieve_bedrag", "Zelfstandigen: alternatieve financiering", "fodsz_z", _BLAD_Z, 35, "103",
         FODSZ_Z, ""),
        ("w_gezinsbijslag", "Werknemers: gezinsbijslag", "fodsz_w", _BLAD_W, 74, "111 KB", FODSZ_W,
         " Verdwijnt vanaf 2015 (6de staatshervorming)."),
        ("w_naar_riziv", "Werknemers: overdrachten naar RIZIV-geneeskundige verzorging", "fodsz_w", _BLAD_W, 85,
         "114.2", FODSZ_W, ""),
        ("w_toelagen_bedrag", "Werknemers: toelagen van publieke overheden", "fodsz_w", _BLAD_W, 23, "102",
         FODSZ_W, ""),
        ("w_alternatieve_bedrag", "Werknemers: alternatieve financiering", "fodsz_w", _BLAD_W, 31, "103",
         FODSZ_W, ""),
        # Samenstelling van de alternatieve financiering (voetnoot bij §6.1): btw en roerende
        # voorheffing; 'andere' = alternatieve financiering - btw - roerende voorheffing.
        ("z_btw_bedrag", "Zelfstandigen: alternatieve financiering uit btw", "fodsz_z", _BLAD_Z, 36,
         "TVA + ART.24§1QUATER", FODSZ_Z, ""),
        ("z_rv_bedrag", "Zelfstandigen: alternatieve financiering uit roerende voorheffing", "fodsz_z", _BLAD_Z,
         43, "RV", FODSZ_Z, ""),
        ("w_btw_bedrag", "Werknemers: alternatieve financiering uit btw", "fodsz_w", _BLAD_W, 32,
         "TVA + ART.24§1QUATER", FODSZ_W, ""),
        ("w_rv_bedrag", "Werknemers: alternatieve financiering uit roerende voorheffing", "fodsz_w", _BLAD_W, 39,
         "RV", FODSZ_W, ""),
    ]
]
