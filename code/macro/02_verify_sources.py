"""
Stap 2: controleren of de overgenomen waarden gelijk zijn aan de originele bron

Voor elke reeks in bronspec.REEKSEN met een `verificatie`-verwijzing leest dit
script het ORIGINELE bronbestand (FOD SZ, Statbel, NBB) en vergelijkt het dat
jaar per jaar met data/raw/*.csv.

Resultaat: output/checks/bronverificatie.csv met per reeks het grootste
verschil en een status (OK / AFWIJKING).

Tolerantie: het werkboek bewaart de consumptieprijsindex afgerond op 2
decimalen, dus daar aanvaarden we een verschil tot 0,005 indexpunt. Voor de
andere reeksen moet het verschil verwaarloosbaar zijn (relatief < 1e-9).

Vereist het pakket `xlrd` om het oude .xls-bestand van Statbel te lezen.
Is dat niet geïnstalleerd, dan slaan we enkel die twee reeksen over.

Uitvoeren:  python code/macro/02_verify_sources.py
"""

import openpyxl
import pandas as pd

from bronspec import REEKSEN
from common import DATA_RAW, DATA_SOURCES, OUT_CHECKS, WERKBOEK, zorg_voor_mappen

# In de FOD SZ-bestanden staat het jaar 2000 in kolom D (= kolom 4).
FODSZ_KOL_2000 = 4
# In het NBB-bestand staat 2003 in kolom D.
NBB_KOL_2003 = 4


def fodsz_rij(blad, rij: int, jaren) -> pd.Series:
    """Lees één rij uit een FOD SZ-werkblad. Een '-' betekent nul."""
    uit = {}
    for j in jaren:
        v = blad.cell(row=rij, column=FODSZ_KOL_2000 + (j - 2000)).value
        uit[j] = 0.0 if v in (None, "-") else float(v)
    return pd.Series(uit)


def statbel_jaargemiddelden(bladnaam: str) -> pd.Series:
    """Haal de jaargemiddelden (basis 2013) uit het historiekbestand van Statbel.

    Het bestand is opgebouwd in blokken van één jaar:
      - een rij met het jaartal,
      - een rij 'Base/Basis' met de basisjaren als kolomkoppen,
      - twaalf maandrijen,
      - een rij 'Jaar/Année' (of 'Année/Jaar') met het jaargemiddelde.
    We zoeken per blok de kolom met basis 2013 en lezen het jaargemiddelde.
    """
    df = pd.read_excel(DATA_SOURCES / "statbel" / "statbel_cpi_historiek_1920-2025.xls",
                       sheet_name=bladnaam, header=None)
    uit = {}
    jaar, kol_2013 = None, None
    for _, rij in df.iterrows():
        waarden = rij.tolist()
        # Een rij met enkel een jaartal (1920-2100) opent een nieuw blok
        getallen = [v for v in waarden if isinstance(v, (int, float)) and not pd.isna(v)]
        tekst = [str(v) for v in waarden if isinstance(v, str)]
        if len(getallen) == 1 and not tekst and 1900 < getallen[0] < 2100 and float(getallen[0]).is_integer():
            jaar = int(getallen[0])
            continue
        if any(t.startswith("Base") for t in tekst):
            kol_2013 = next((i for i, v in enumerate(waarden) if v == 2013 or v == 2013.0), None)
            continue
        if any(t in ("Jaar/Année", "Année/Jaar") for t in tekst) and jaar and kol_2013 is not None:
            uit[jaar] = float(waarden[kol_2013])
    return pd.Series(uit)


def main() -> None:
    zorg_voor_mappen()

    # Bronbestanden openen
    fz = openpyxl.load_workbook(DATA_SOURCES / "fodsz" / "fodsz_zelfstandigen_2000-2025.xlsm", data_only=True)
    fw = openpyxl.load_workbook(DATA_SOURCES / "fodsz" / "fodsz_werknemers_2000-2025.xlsm", data_only=True)
    blad_z, blad_w = fz.worksheets[0], fw.worksheets[0]
    nbb = openpyxl.load_workbook(DATA_SOURCES / "nbb" / "nbb_sociale_premies_2003-2023.xlsx", data_only=True)["Table"]
    willem = openpyxl.load_workbook(WERKBOEK, data_only=True)

    # Alle overgenomen reeksen in één tabel
    raw = {}
    for pad in DATA_RAW.glob("*.csv"):
        if pad.name == "codebook.csv":
            continue
        # data/raw/ bevat ook micro-invoer (abc_*, kakwani_*, micro_*) zonder kolom 'jaar':
        # die hoort niet bij de macroreeksen en wordt hier overgeslagen
        if "jaar" not in pd.read_csv(pad, nrows=0).columns:
            continue
        raw[pad.stem] = pd.read_csv(pad, index_col="jaar")

    statbel_cache = {}
    resultaten = []

    for r in REEKSEN:
        if r.verificatie is None:
            status = ("RECHTSTREEKS UIT BRON (geen controle nodig)" if r.werkboek.startswith("fodsz")
                      else "NIET GEVERIFIEERD (geen bronbestand)")
            resultaten.append({"variabele": r.variabele, "bron": r.bron_bestand, "max_abs_verschil": None,
                               "max_rel_verschil": None, "status": status})
            continue

        jaren = list(range(r.jaren[0], r.jaren[1] + 1))
        overgenomen = raw[r.csv][r.variabele].loc[jaren]
        soort = r.verificatie[0]
        tolerantie_abs = 0.0

        if soort in ("fodsz_z", "fodsz_w"):
            blad = blad_z if soort == "fodsz_z" else blad_w
            bron = fodsz_rij(blad, r.verificatie[1], jaren)

        elif soort in ("fodsz_z_aandeel", "fodsz_w_aandeel"):
            # Aandeel = som van teller-rij(en) / totaal lopende ontvangsten
            blad = blad_z if soort.startswith("fodsz_z") else blad_w
            tellers = r.verificatie[1] if isinstance(r.verificatie[1], tuple) else (r.verificatie[1],)
            teller = sum(fodsz_rij(blad, t, jaren) for t in tellers)
            bron = teller / fodsz_rij(blad, r.verificatie[2], jaren)

        elif soort == "nbb":
            bron = pd.Series({j: float(nbb.cell(row=r.verificatie[1], column=NBB_KOL_2003 + (j - 2003)).value)
                              for j in jaren})

        elif soort == "statbel":
            try:
                if r.verificatie[1] not in statbel_cache:
                    statbel_cache[r.verificatie[1]] = statbel_jaargemiddelden(r.verificatie[1])
            except ImportError:
                resultaten.append({"variabele": r.variabele, "bron": r.bron_bestand, "status":
                                   "OVERGESLAGEN (installeer xlrd: python -m pip install --user xlrd)"})
                continue
            bron = statbel_cache[r.verificatie[1]].loc[jaren]
            # Het werkboek rondt de CPI af op 2 decimalen
            tolerantie_abs = 0.00501 if r.variabele == "cpi_2013" else 1e-6

        elif soort == "willem_rij":
            # Vergelijking met een tweede plaats in het werkboek (enkel 2003-2023 aanwezig)
            blad = willem[r.verificatie[1]]
            jaren = [j for j in jaren if 2003 <= j <= 2023]
            overgenomen = overgenomen.loc[jaren]
            bron = pd.Series({j: float(blad.cell(row=r.verificatie[2], column=2 + (j - 2003)).value) for j in jaren})

        else:
            raise ValueError(f"onbekende verificatie: {soort}")

        verschil = (overgenomen - bron).abs()
        rel = (verschil / bron.abs().where(bron != 0)).fillna(0)
        ok = bool(((verschil <= tolerantie_abs) | (rel < 1e-9)).all())
        resultaten.append({
            "variabele": r.variabele,
            "bron": r.bron_bestand,
            "jaren_gecontroleerd": f"{jaren[0]}-{jaren[-1]}",
            "max_abs_verschil": verschil.max(),
            "max_rel_verschil": rel.max(),
            "tolerantie_abs": tolerantie_abs,
            "status": "OK" if ok else "AFWIJKING",
            "jaren_met_afwijking": ", ".join(str(j) for j in verschil.index
                                             if not (verschil[j] <= tolerantie_abs or rel[j] < 1e-9)),
        })

    df = pd.DataFrame(resultaten)
    df.to_csv(OUT_CHECKS / "bronverificatie.csv", index=False)
    print(df[["variabele", "status", "max_abs_verschil", "jaren_met_afwijking"]].to_string(index=False))
    n_afw = (df["status"] == "AFWIJKING").sum()
    print(f"\n{n_afw} reeks(en) met afwijking. Details: output/checks/bronverificatie.csv")


if __name__ == "__main__":
    main()
